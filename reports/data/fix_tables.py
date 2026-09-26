"""
Session 7b — makes every table fit the text column and read cleanly.

Diagnosis
---------
1. OVERFLOW. Word recomputes `w:tblGrid` from cell content whenever the table
   layout is "autofit". Table 5.2's grid summed to 10,636 twips against a text
   width of 9,026, so its last column ran off the right edge of the page.
   Setting `w:tblW` alone does not prevent this; the layout must be FIXED and
   the grid written explicitly.

2. JUSTIFIED CELL TEXT. The `Normal` style in this document is JUSTIFY. Inside a
   narrow cell that stretches a two-word line to both edges, which is why
   "astropy__astropy-7746" rendered as a right-aligned fragment. Cell paragraphs
   now set LEFT explicitly.

3. CAPTION SEPARATED FROM TABLE. A caption with no "keep with next" is left
   stranded at the bottom of a page while its table moves to the next one.

Column widths are proportional to content rather than equal, so narrow columns
("Class", "Round 1") stay narrow and prose columns get the space, subject to a
minimum readable width.

Run from the project root:
    python3 reports/data/fix_tables.py --in-place
"""

import sys
import os
import re

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import docx
from docx.shared import Pt, Twips
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from docx_writer import set_table_borders, backup  # noqa: E402

# Layout note: these scripts live in reports/data/, so the project root is
# two levels up. REPORTS anchors figures/, which moved alongside them.
REPORTS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(REPORTS)
SRC = os.path.join(ROOT, "Self.docx")

TEXT_W = 9026        # twips available between the margins
MIN_COL = 640        # no column narrower than ~0.44 in
CAPTION = re.compile(r'^Table \d+\.\d+')


def col_weights(table):
    """Weight each column by its content, blending the longest and mean cell."""
    ncols = len(table.columns)
    longest = [1] * ncols
    means = [1.0] * ncols
    for j in range(ncols):
        lens = []
        for row in table.rows:
            try:
                txt = row.cells[j].text.strip()
            except IndexError:
                continue
            # the longest single word sets the minimum sensible column width
            word = max((len(w) for w in txt.split()), default=1)
            lens.append(max(len(txt), word * 1.6))
        if lens:
            longest[j] = max(lens)
            means[j] = sum(lens) / len(lens)
    # a column's demand is mostly its worst case, tempered by its average
    return [0.65 * longest[j] + 0.35 * means[j] for j in range(ncols)]


def allocate(weights, total=TEXT_W, min_col=MIN_COL):
    """Split `total` across columns proportionally, respecting a floor."""
    n = len(weights)
    if n * min_col >= total:
        base = total // n
        out = [base] * n
        out[-1] += total - base * n
        return out
    widths = [max(min_col, int(total * w / sum(weights))) for w in weights]
    # iteratively pull the excess out of the columns that are above the floor
    for _ in range(40):
        diff = total - sum(widths)
        if diff == 0:
            break
        flexible = [i for i, w in enumerate(widths) if w > min_col]
        if not flexible:
            widths[-1] += diff
            break
        share = diff / len(flexible)
        for i in flexible:
            widths[i] = max(min_col, int(widths[i] + share))
    widths[-1] += total - sum(widths)
    return widths


def set_grid(table, widths):
    tbl = table._tbl
    old = tbl.find(qn('w:tblGrid'))
    if old is not None:
        tbl.remove(old)
    grid = OxmlElement('w:tblGrid')
    for w in widths:
        gc = OxmlElement('w:gridCol')
        gc.set(qn('w:w'), str(int(w)))
        grid.append(gc)
    # tblGrid must sit immediately after tblPr
    pr = tbl.tblPr
    pr.addnext(grid)


def set_prop(tblPr, tag, attrs):
    old = tblPr.find(qn(tag))
    if old is not None:
        tblPr.remove(old)
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(k), v)
    tblPr.append(el)
    return el


def fix_table(table, font_pt):
    ncols = len(table.columns)
    widths = allocate(col_weights(table))
    tblPr = table._tbl.tblPr

    table.autofit = False
    set_prop(tblPr, 'w:tblW', {'w:w': str(TEXT_W), 'w:type': 'dxa'})
    set_prop(tblPr, 'w:tblLayout', {'w:type': 'fixed'})
    set_grid(table, widths)
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row in table.rows:
        for j, cell in enumerate(row.cells):
            if j < len(widths):
                cell.width = Twips(int(widths[j]))
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf = p.paragraph_format
                pf.space_before = Pt(1)
                pf.space_after = Pt(1)
                pf.left_indent = Pt(0)
                pf.first_line_indent = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(font_pt)
    return widths


def main():
    in_place = "--in-place" in sys.argv
    doc = docx.Document(SRC)

    # keep each table caption on the same page as its table
    kept = 0
    for p in doc.paragraphs:
        if CAPTION.match(p.text.strip()):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            kept += 1
    print(f"keep-with-next set on {kept} table captions\n")

    for i, t in enumerate(doc.tables):
        ncols = len(t.columns)
        font = 9 if ncols >= 6 else 10 if ncols >= 4 else 11
        w = fix_table(t, font)
        print(f"  table[{i:2}] {len(t.rows):>2}r x {ncols}c  font {font}pt  "
              f"cols={w}  sum={sum(w)}")

    if in_place:
        b = backup(SRC)
        doc.save(SRC)
        print(f"\nUpdated Self.docx; backup {os.path.basename(b)}")
    else:
        doc.save(os.path.join(ROOT, "Self_tab.docx"))
        print("\nWrote Self_tab.docx")


if __name__ == "__main__":
    main()
