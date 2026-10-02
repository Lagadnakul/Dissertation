"""
Session 7 — final formatting consistency pass over Self.docx.

Defects found by the whole-document audit:

  A. Chapter 3 used 16 pt for both "CHAPTER 3" and "LITERATURE REVIEW", where
     Chapters 1, 2, 4, 5 and 6 use 18 pt for the chapter line and 14 pt for the
     title line.
  B. Four section headings deviated from the 14 pt used by the other 65:
     "1.7 Problem Statement" was 16 pt, and "1.10", "2.6" and "2.7" carried no
     explicit size so they rendered at the 12 pt body size.
  C. The two front-matter tables (approval-certificate signature block and the
     abbreviations list) had no explicit table borders, relying on the table
     style alone, unlike every other table in the document.

Run from the project root:
    python3 reports/data/fix_format.py --in-place
"""

import sys
import os

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_writer import set_table_borders, fit_table_width, backup  # noqa: E402

# Layout note: these scripts live in reports/data/, so the project root is
# two levels up. REPORTS anchors figures/, which moved alongside them.
REPORTS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(REPORTS)
SRC = os.path.join(ROOT, "Self.docx")

CHAPTER_PT, SECTION_PT = 18, 14


def resize(p, pt, centre=None, bold=True):
    if not p.runs:
        return False
    for r in p.runs:
        r.font.size = Pt(pt)
        if bold:
            r.font.bold = True
    if centre is not None:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centre else None
    return True


def main():
    in_place = "--in-place" in sys.argv
    doc = docx.Document(SRC)
    ps = doc.paragraphs

    print("A. CHAPTER HEADINGS")
    for i, p in enumerate(ps):
        if p.text.strip() == "CHAPTER 3":
            resize(p, CHAPTER_PT, centre=True)
            resize(ps[i + 1], SECTION_PT, centre=True)
            print(f"   CHAPTER 3 -> {CHAPTER_PT} pt, "
                  f"{ps[i+1].text.strip()} -> {SECTION_PT} pt")

    print("\nB. SECTION HEADINGS")
    for target in ("1.7 Problem Statement", "1.10 Scope of the Dissertation",
                   "2.6 Limitations of the Study", "2.7 Chapter Summary"):
        for p in ps:
            if p.text.strip() == target:
                before = (p.runs[0].font.size.pt
                          if p.runs and p.runs[0].font.size else "inherited")
                resize(p, SECTION_PT)
                print(f"   {target:<34} {before} -> {SECTION_PT} pt")
                break

    print("\nC. FRONT-MATTER TABLES")
    for idx in (0, 1):
        t = doc.tables[idx]
        set_table_borders(t)
        fit_table_width(t)
        print(f"   table[{idx}] borders added, width normalised to 9024 twips")

    if in_place:
        b = backup(SRC)
        doc.save(SRC)
        print(f"\nUpdated Self.docx; backup {os.path.basename(b)}")
    else:
        doc.save(os.path.join(ROOT, "Self_fmt.docx"))
        print("\nWrote Self_fmt.docx")


if __name__ == "__main__":
    main()
