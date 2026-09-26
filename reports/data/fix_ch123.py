"""
Session 7c — brings Chapters 1 to 3 into line with the conventions used in
Chapters 4 to 6.

Defects found by the whole-document audit:

  1. BOLD BODY TEXT. 27 full sentences in Chapters 1-3 are set in bold.
     Chapters 4-6 contain none. Three sub-cases are treated differently:
       a. A label followed by prose ("Reasoning correctness - whether the agent
          ...", "Chapter 3 - Literature Review presents ...", "Blind Retry, in
          which ...") keeps the label bold and drops the bold from the prose.
       b. A research question sitting under an already-bold label
          ("RQ1. Recovery Effectiveness") loses its bold entirely.
       c. Short labelled statements ("Objective 4: ...", "Assumption 1: ...")
          are genuine labels and keep their bold.

  2. CAPTION POSITION. Figures 1.1, 1.2, 1.3, 3.1 and 3.3 carry their caption
     ABOVE the image, while Figure 3.2 and every figure in Chapters 4-6 carry it
     below. Captions are moved below the image throughout, which is both the
     standard convention and what the majority of the document already does.

  3. FIGURE ALIGNMENT. Figure images and their captions are centred document
     wide, and the image paragraph keeps with the caption so the two cannot be
     split across a page break.

Image sizes in Chapters 1-3 are deliberately left alone. They range from 4.33 to
5.84 inches; enlarging them to the 6.00 inch width used in Chapters 4-6 would
upscale existing raster images and visibly blur them.

Run from the project root:
    python3 reports/data/fix_ch123.py --in-place
"""

import sys
import os
import re

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import docx
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_writer import backup  # noqa: E402

# Layout note: these scripts live in reports/data/, so the project root is
# two levels up. REPORTS anchors figures/, which moved alongside them.
REPORTS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(REPORTS)
SRC = os.path.join(ROOT, "Self.docx")

# (a) label + prose: bold only the prefix
SPLIT_AFTER = [
    "Reasoning correctness",
    "Patch/application reliability",
    "Chapter 2 — Aim, Objectives and Scope",
    "Chapter 3 — Literature Review",
    "Chapter 4 — Methodology",
    "Chapter 5 — Results and Discussion",
    "Chapter 6 — Conclusions and Future Work",
    "Blind Retry",
    "Reflection-only",
    "Diagnose+Revise",
]

# (b) whole paragraph loses its bold
UNBOLD_PREFIX = [
    "Does diagnosis-based recovery improve",
    "Does providing the agent with its own previous failed attempt",
    "To what extent do patch-formatting and patch-application failures",
    "Are the observed recovery outcomes consistent across independent",
    "There is limited empirical evidence on whether execution-grounded",
]

CAPTION = re.compile(r'^Figure \d+\.\d+\s')
# captions that sit above their image and must move below it
MOVE_BELOW = ["Figure 1.1", "Figure 1.2", "Figure 1.3",
              "Figure 3.1", "Figure 3.3"]


def has_img(p):
    return bool(p._p.findall('.//' + qn('a:blip')))


def set_plain(p, text):
    """Rewrite a paragraph as a single non-bold run."""
    first = p.runs[0]
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)
    first.text = text
    first.font.bold = None


def split_bold(p, label):
    """Bold `label`, leave the remainder plain, inside one paragraph."""
    full = p.text
    rest = full[len(label):]
    first = p.runs[0]
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)
    first.text = label
    first.font.bold = True
    tail = p.add_run(rest)
    tail.font.bold = None
    tail.font.size = first.font.size
    return True


def main():
    in_place = "--in-place" in sys.argv
    doc = docx.Document(SRC)
    ps = doc.paragraphs
    ch4 = next(i for i, p in enumerate(ps) if p.text.strip() == "CHAPTER 4")

    print("1. BOLD BODY TEXT")
    split_n = unbold_n = 0
    for i in range(ch4):
        p = ps[i]
        t = p.text.strip()
        if not t or not p.runs or not p.runs[0].font.bold or len(t) <= 90:
            continue
        for lab in SPLIT_AFTER:
            if t.startswith(lab) and len(t) > len(lab) + 5:
                split_bold(p, lab)
                split_n += 1
                print(f"   para {i:4} bold kept on label only: {lab!r}")
                break
        else:
            for pre in UNBOLD_PREFIX:
                if t.startswith(pre):
                    set_plain(p, p.text)
                    unbold_n += 1
                    print(f"   para {i:4} un-bolded: {t[:62]}...")
                    break
    print(f"   -> {split_n} split, {unbold_n} un-bolded\n")

    print("2. CAPTION POSITION")
    moved = 0
    for label in MOVE_BELOW:
        ps = doc.paragraphs
        ci = next((i for i, p in enumerate(ps)
                   if p.text.strip().startswith(label)), None)
        if ci is None:
            print(f"   {label}: caption not found")
            continue
        # find the run of image paragraphs that follows the caption
        j = ci + 1
        last_img = None
        while j < len(ps) and (has_img(ps[j]) or not ps[j].text.strip()):
            if has_img(ps[j]):
                last_img = j
            j += 1
            if j - ci > 6:
                break
        if last_img is None:
            print(f"   {label}: no image below caption, left as is")
            continue
        ps[last_img]._p.addnext(ps[ci]._p)   # move caption after the image
        moved += 1
        print(f"   {label}: caption {ci} moved below image {last_img}")
    print(f"   -> {moved} captions moved\n")

    print("3. FIGURE ALIGNMENT")
    ps = doc.paragraphs
    cap_n = img_n = 0
    for i, p in enumerate(ps):
        if CAPTION.match(p.text.strip()) and len(p.text.strip()) < 110:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(10)
            cap_n += 1
        if has_img(p):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            img_n += 1
    print(f"   {cap_n} captions centred, {img_n} images centred and kept with caption")

    if in_place:
        b = backup(SRC)
        doc.save(SRC)
        print(f"\nUpdated Self.docx; backup {os.path.basename(b)}")
    else:
        doc.save(os.path.join(ROOT, "Self_c13.docx"))
        print("\nWrote Self_c13.docx")


if __name__ == "__main__":
    main()
