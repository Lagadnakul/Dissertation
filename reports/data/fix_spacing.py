"""
Session 7e — fixes the excessive vertical spacing and the resulting blank pages.

Diagnosis
---------
The `Normal` style carries space_after = 19 pt and line_spacing = 1.467, and
every paragraph in the document uses `Normal`. The consequences:

  * A one-word bullet such as "planning failure;" occupies 19 pt of trailing
    space plus a 1.47-spaced line -- roughly two and a half times its own
    height. Section 3.10 and the Section 3.11 "positioning ladder" are built
    from many such short paragraphs, which is why they sprawl over whole pages.

  * Every paragraph INSIDE a table cell inherits the same 1.467 line spacing.
    Table 3.1 is 36 rows by 6 columns with rows up to six lines deep, so the
    inherited spacing added roughly 47 per cent to the height of the largest
    table in the document, spreading it over far more pages than it needs and
    leaving near-empty pages around it.

Fix
---
  1. `Normal`: space_after 19 pt -> 6 pt, line_spacing set to exactly 1.5
     (the usual dissertation setting; 1.467 was an artefact).
  2. Table cells: line_spacing 1.0 and 2 pt of padding, set explicitly so they
     no longer inherit body spacing.
  3. Table rows are allowed to break across pages, so a long table flows instead
     of being pushed whole onto the next page.

Run from the project root:
    python3 reports/data/fix_spacing.py --in-place
"""

import sys
import os

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import docx
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from docx_writer import backup  # noqa: E402

# Layout note: these scripts live in reports/data/, so the project root is
# two levels up. REPORTS anchors figures/, which moved alongside them.
REPORTS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(REPORTS)
SRC = os.path.join(ROOT, "Self.docx")

BODY_AFTER = 6.0
BODY_LINE = 1.5


def allow_row_split(table):
    """Clear w:cantSplit so long tables flow across pages."""
    n = 0
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        cs = trPr.find(qn('w:cantSplit'))
        if cs is not None:
            trPr.remove(cs)
            n += 1
    return n


def repeat_header(table):
    """Mark the first row as a repeating header row."""
    trPr = table.rows[0]._tr.get_or_add_trPr()
    if trPr.find(qn('w:tblHeader')) is None:
        trPr.append(OxmlElement('w:tblHeader'))


def main():
    in_place = "--in-place" in sys.argv
    doc = docx.Document(SRC)

    print("1. NORMAL STYLE")
    pf = doc.styles['Normal'].paragraph_format
    before = (pf.space_after.pt if pf.space_after else None, pf.line_spacing)
    pf.space_after = Pt(BODY_AFTER)
    pf.line_spacing = BODY_LINE
    print(f"   space_after {before[0]} pt -> {BODY_AFTER} pt")
    print(f"   line_spacing {before[1]:.3f} -> {BODY_LINE}")

    print("\n2. TABLE CELL SPACING")
    cells = 0
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    cpf = p.paragraph_format
                    cpf.line_spacing = 1.0
                    cpf.space_before = Pt(2)
                    cpf.space_after = Pt(2)
                    cells += 1
    print(f"   {cells} cell paragraphs set to single spacing, 2 pt padding")

    print("\n3. LONG TABLES FLOW ACROSS PAGES")
    for i, t in enumerate(doc.tables):
        cleared = allow_row_split(t)
        if len(t.rows) > 8:
            repeat_header(t)
            print(f"   table[{i:2}] {len(t.rows):>2} rows: header row repeats, "
                  f"{cleared} cantSplit cleared")

    # paragraphs that explicitly carried the old 19 pt-ish spacing
    print("\n4. STRAY EXPLICIT SPACING ON BODY PARAGRAPHS")
    fixed = 0
    for p in doc.paragraphs:
        sp = p.paragraph_format.space_after
        if sp is not None and sp.pt > 12:
            p.paragraph_format.space_after = Pt(BODY_AFTER)
            fixed += 1
    print(f"   {fixed} paragraphs with space_after > 12 pt normalised")

    if in_place:
        b = backup(SRC)
        doc.save(SRC)
        print(f"\nUpdated Self.docx; backup {os.path.basename(b)}")
    else:
        doc.save(os.path.join(ROOT, "Self_sp.docx"))
        print("\nWrote Self_sp.docx")


if __name__ == "__main__":
    main()
