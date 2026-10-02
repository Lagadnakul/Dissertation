"""
Session 7d — adds page numbers to the footer.

The document had no footers and no PAGE field anywhere, yet its Table of
Contents, List of Figures and List of Tables all cite page numbers. Without a
printed number on the page those references cannot be followed.

A centred PAGE field is inserted into the footer of every section. Word
evaluates it on open, so the numbers appear immediately and stay correct.

Run from the project root:
    python3 reports/data/add_pagenumbers.py --in-place
"""

import sys
import os

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import docx
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_writer import backup  # noqa: E402

# Layout note: these scripts live in reports/data/, so the project root is
# two levels up. REPORTS anchors figures/, which moved alongside them.
REPORTS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(REPORTS)
SRC = os.path.join(ROOT, "Self.docx")


def add_page_field(paragraph):
    """Insert a { PAGE } field so Word renders the live page number."""
    run = paragraph.add_run()
    r = run._r

    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE   \\* MERGEFORMAT '
    sep = OxmlElement('w:fldChar')
    sep.set(qn('w:fldCharType'), 'separate')
    txt = OxmlElement('w:t')
    txt.text = "1"                      # cached value, refreshed by Word
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')

    for el in (begin, instr, sep, txt, end):
        r.append(el)
    run.font.size = Pt(11)
    return run


def main():
    in_place = "--in-place" in sys.argv
    doc = docx.Document(SRC)

    for i, section in enumerate(doc.sections):
        section.footer.is_linked_to_previous = False
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_field(p)
        print(f"  section {i}: centred PAGE field added to footer")

    if in_place:
        b = backup(SRC)
        doc.save(SRC)
        print(f"\nUpdated Self.docx; backup {os.path.basename(b)}")
    else:
        doc.save(os.path.join(ROOT, "Self_pn.docx"))
        print("\nWrote Self_pn.docx")


if __name__ == "__main__":
    main()
