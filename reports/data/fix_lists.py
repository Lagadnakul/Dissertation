"""
Session 7 — repairs the Chapter 3 figure captions and rebuilds the
List of Figures and List of Tables.

Defects addressed (PLAN.md section D9):
  D9-1   Both lists were carried over from an unrelated thesis. Every row
         referred to Split CIFAR-100 and a meta-continual learning framework.
         Rebuilt from the captions actually present in this document.
  D9-10  Figure 3.2 exists as an image in Section 3.6 but had no caption and was
         referenced nowhere. Caption added.
  Fig 3.1 caption carried editing scaffolding ("-> Section 3.2") and used an
         em dash where Chapters 1 and 4-6 use none. Normalised.

Page numbers are deliberately left blank: they depend on how Word paginates the
images and tables, which cannot be determined without rendering the document.
Nakul fills these once, in Word, after final pagination.

Run from the project root:
    python3 reports/data/fix_lists.py             # writes Self_lists.docx
    python3 reports/data/fix_lists.py --in-place  # updates Self.docx, backs up first
"""

import sys
import os
import re

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import docx
from docx_writer import set_table_borders, fit_table_width, backup  # noqa: E402

# Layout note: these scripts live in reports/data/, so the project root is
# two levels up. REPORTS anchors figures/, which moved alongside them.
REPORTS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(REPORTS)
SRC = os.path.join(ROOT, "Self.docx")

CAPTION_FIXES = [
    ("Figure 3.1 — General Workflow of an Autonomous Coding Agent → Section 3.2",
     "Figure 3.1 General Workflow of an Autonomous Coding Agent"),
    ("Figure 3.3 — Positioning and Recovery Strategy Framework of the Present Study",
     "Figure 3.3 Positioning and Recovery Strategy Framework of the Present Study"),
]

FIG_32_CAPTION = "Figure 3.2 Failure Categories in Repository-Level Coding"


def set_cell(cell, text, bold=False):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    run.font.bold = bold or None


def rebuild(table, header, rows):
    """Resize `table` to fit `rows` and write them in."""
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    for i, h in enumerate(header):
        set_cell(table.rows[0].cells[i], h, bold=True)
    for r in rows:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            set_cell(cells[i], v)
    set_table_borders(table)
    fit_table_width(table)


def main():
    in_place = "--in-place" in sys.argv
    doc = docx.Document(SRC)
    ps = doc.paragraphs

    # ---------------------------------------------- caption repairs
    fixed = 0
    for p in ps:
        t = p.text.strip()
        for old, new in CAPTION_FIXES:
            if t == old:
                for r in list(p.runs):
                    r._element.getparent().remove(r._element)
                p.add_run(new)
                print(f"  caption fixed: {new}")
                fixed += 1

    # Figure 3.2: insert a caption after its image, before heading 3.7
    if not any(p.text.strip().startswith("Figure 3.2") for p in ps):
        anchor = next(p for p in ps
                      if p.text.strip() == "3.7 Automated Debugging and Program Repair")
        cap = anchor.insert_paragraph_before("")
        cap.style = doc.styles["Normal"]
        cap.add_run(FIG_32_CAPTION)
        print(f"  caption added: {FIG_32_CAPTION}")
        fixed += 1

    ps = doc.paragraphs  # refresh after insertion

    # ---------------------------------------------- collect captions
    figs, tabs = [], []
    body_start = 205  # first paragraph of Chapter 1
    for i, p in enumerate(ps):
        t = p.text.strip()
        if i < body_start or len(t) > 120:
            continue
        m = re.match(r'^Figure (\d+\.\d+)\s+(.+)$', t)
        if m:
            figs.append((m.group(1), m.group(2).strip()))
            continue
        m = re.match(r'^Table (\d+\.\d+):?\s+(.+)$', t)
        if m:
            tabs.append((m.group(1), m.group(2).strip()))

    def key(x):
        return [int(v) for v in x[0].split('.')]
    figs.sort(key=key)
    tabs.sort(key=key)

    # ---------------------------------------------- rebuild the two lists
    rebuild(doc.tables[2], ["Figure No.", "Title", "Page Number"],
            [(n, t, "") for n, t in figs])
    rebuild(doc.tables[3], ["Table No.", "Title", "Page Number"],
            [(n, t, "") for n, t in tabs])

    print(f"\n  List of Figures: {len(figs)} entries")
    for n, t in figs:
        print(f"    {n:<6} {t}")
    print(f"\n  List of Tables: {len(tabs)} entries")
    for n, t in tabs:
        print(f"    {n:<6} {t}")

    if in_place:
        b = backup(SRC)
        doc.save(SRC)
        print(f"\nUpdated Self.docx; backup {os.path.basename(b)}")
    else:
        out = os.path.join(ROOT, "Self_lists.docx")
        doc.save(out)
        print(f"\nWrote {out}")


if __name__ == "__main__":
    main()
